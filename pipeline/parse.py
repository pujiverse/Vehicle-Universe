import openpyxl,glob,re,json,os
U='/root/.claude/uploads/3e5d0190-f1ae-56ea-907f-0e19d9663d06/'
recs=[]; brands={}
for f in sorted(glob.glob(U+'*.xlsx')):
    m=re.search(r'-(\d\d)_',os.path.basename(f)); n=int(m.group(1))
    if n==0: continue
    t='VT%02d'%n
    wb=openpyxl.load_workbook(f)
    bs=wb['Brands']
    binfo={}
    for r in bs.iter_rows(min_row=5,values_only=True):
        if r[1]: binfo[r[1]]={'country':r[2],'website':r[3] if r[3] not in ('-',None,'not recorded') else None,'classes':r[8]}
    for name in wb.sheetnames[1:]:
        ws=wb[name]
        title=ws.cell(1,1).value or ''
        bname=title.rsplit(' - ',1)[0] if ' - ' in title else name
        bi=binfo.get(bname) or binfo.get(name) or {}
        for row in ws.iter_rows(min_row=5):
            v=[c.value for c in row]
            if not v or not v[0]: continue
            link=row[9].hyperlink.target if len(row)>9 and row[9].hyperlink else None
            q=re.search(r'(Q\d+)',link or '')
            recs.append({'type':t,'brand':bname,'country':bi.get('country'),'brand_site':bi.get('website'),'model':str(v[0]),'desc':v[1],'year':v[2],'basis':v[3],'dates':v[4],'price':v[5],'unit':v[6],'mfrs':v[7],'model_site':v[8] if v[8] not in ('-',None) else None,'qid':q.group(1) if q else None})
json.dump(recs,open('recs.json','w'),ensure_ascii=False)
import collections
print(len(recs),'records', sum(1 for r in recs if r['qid']),'with qid', len({r['qid'] for r in recs}),'unique qid')
print(collections.Counter(r['type'] for r in recs))
print('prices',sum(1 for r in recs if r['price'] is not None), collections.Counter(r['unit'] for r in recs if r['price'] is not None).most_common(10))
print('model sites',sum(1 for r in recs if r['model_site'])); print('years',sum(1 for r in recs if isinstance(r['year'],int)))
print(collections.Counter(r['basis'] for r in recs).most_common(10))
print(recs[100]); print(len({r['brand'] for r in recs}),'brands')
